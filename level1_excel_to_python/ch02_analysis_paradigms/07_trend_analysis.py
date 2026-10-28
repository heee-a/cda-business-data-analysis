#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
实验 07：趋势分析（对应 2、由基础分析范式引申出的分析方法 → 1 趋势分析）
==================================================================
【原 Excel/Power BI 实现内涵】把时间序列画成折线，观察长期走势、拐点与周期。
【Python 迁移要点】用 openpyxl 在含多张表的模板里定位“year/订单量”表，
读取出后画折线并叠加 3 期移动平均以凸显趋势。
"""
import os
import sys
import openpyxl
import pandas as pd
import matplotlib.pyplot as plt

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from common import DATA_LEVEL1, save_fig, section

SRC = os.path.join(DATA_LEVEL1, "trend", "1趋势分析.xlsx")
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "output", "trend")


def main():
    section("1. 定位并读取 year / 订单量 表")
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb.active
    year_col = val_col = header_row = None
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        for j, v in enumerate(row):
            if isinstance(v, str) and v.strip() == "year":
                year_col, val_col, header_row = j, j + 1, i
                break
        if header_row:
            break
    rows = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if row[year_col] is None:          # 遇到空行即停止（避免读到下方的其他表）
            break
        if isinstance(row[year_col], (int, float)) and isinstance(row[val_col], (int, float)):
            rows.append((int(row[year_col]), row[val_col]))
    df = pd.DataFrame(rows, columns=["year", "订单量"])
    print(df)

    section("2. 趋势 + 3 期移动平均")
    df["移动平均"] = df["订单量"].rolling(3, min_periods=1).mean().round(1)
    fig, ax = plt.subplots(figsize=(10, 5))
    ax.plot(df["year"], df["订单量"], marker="o", color="#4C72B0", label="订单量")
    ax.plot(df["year"], df["移动平均"], color="#d62728", lw=2, label="3 期移动平均")
    ax.set_title("近 10 年订单量趋势")
    ax.set_xlabel("年份"); ax.set_ylabel("订单量（个）"); ax.legend()
    path = save_fig(fig, "trend.png", OUT)
    print("已保存：", path)


if __name__ == "__main__":
    main()
