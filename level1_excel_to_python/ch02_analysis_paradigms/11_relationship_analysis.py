#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
实验 11：关系分析（对应 2、引申出的分析方法 → 5 关系分析）
==================================================================
【原 Excel/Power BI 实现内涵】关系分析既看“分类变量之间”(如不同营销活动
的下单/未下单人数)，也看“两个连续变量之间”(投入成本 vs 产出收益)。教材把
这两类关系放在一张表里。
【Python 迁移要点】用 openpyxl 取出两块：①各营销活动的下单/未下单对比(分组条形)；
②产出收益~投入成本(散点+相关系数)。通过表头名称定位列，避免多表模板干扰。
"""
import os
import sys
import openpyxl
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from scipy import stats

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from common import DATA_LEVEL1, save_fig, section

SRC = os.path.join(DATA_LEVEL1, "relationship", "5关系分析.xlsx")
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "output", "relationship")


def main():
    section("1. 读取两块区域")
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb.active
    # 用第 1 行建立 表头名 -> 列号 映射
    header = {}
    first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    for j, v in enumerate(first):
        if isinstance(v, str):
            header[v.strip()] = j
    rev_col = header.get("产出收益")
    cost_col = header.get("投入成本")
    activities, cost, revenue = [], [], []
    for row in ws.iter_rows(values_only=True):
        b = row[1]  # B 列标签
        if isinstance(b, str) and b.startswith("营销活动"):
            activities.append((b, row[2], row[3]))
        if rev_col is not None and cost_col is not None:
            rv, cv = row[rev_col], row[cost_col]
            if isinstance(rv, (int, float)) and isinstance(cv, (int, float)):
                revenue.append(rv); cost.append(cv)
    act = pd.DataFrame(activities, columns=["营销活动", "下单用户", "未下单用户"])
    pair = pd.DataFrame({"投入成本": cost, "产出收益": revenue})
    print(act)
    print(pair.head())

    section("2. 分类变量关系：各营销活动下单 vs 未下单")
    fig, ax = plt.subplots(1, 2, figsize=(14, 5))
    x = np.arange(len(act)); w = 0.4
    ax[0].bar(x - w/2, act["下单用户"], w, label="下单用户", color="#4C72B0")
    ax[0].bar(x + w/2, act["未下单用户"], w, label="未下单用户", color="#dd8452")
    ax[0].set_xticks(x); ax[0].set_xticklabels(act["营销活动"])
    ax[0].set_title("不同营销活动的下单/未下单人数"); ax[0].legend()

    section("3. 连续变量关系：投入成本 vs 产出收益")
    r, p = stats.pearsonr(pair["投入成本"], pair["产出收益"])
    ax[1].scatter(pair["投入成本"], pair["产出收益"], alpha=0.6, color="#55a868")
    b1, b0 = np.polyfit(pair["投入成本"], pair["产出收益"], 1)
    xs = np.linspace(pair["投入成本"].min(), pair["投入成本"].max(), 50)
    ax[1].plot(xs, b0 + b1 * xs, color="#d62728", lw=2)
    ax[1].set_title(f"投入成本 vs 产出收益 (r={r:.3f})")
    ax[1].set_xlabel("投入成本"); ax[1].set_ylabel("产出收益")
    path = save_fig(fig, "relationship.png", OUT)
    print("已保存：", path)
    print(f"Pearson r={r:.4f}, p={p:.2e} → 投入成本与产出收益呈"
          f"{'正' if r>0 else '负'}相关。")


if __name__ == "__main__":
    main()
